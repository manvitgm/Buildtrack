"""
Generate Improved Material Management System Presentation
Matches the original template: white background, teal header, dark text.
Includes DFD & ER diagram images from the project report.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import os

# ── Colors matching original template ────────────────────────────────────────
BG_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)       # slide background (content slides)
TITLE_BG  = RGBColor(0x02, 0x80, 0x90)       # title slide bg & header bar
HEADER_BG = RGBColor(0x02, 0x80, 0x90)       # header bar
ACCENT    = RGBColor(0x02, 0xC3, 0x9A)        # teal accent (subtitle, bullets)
FOOTER_BG = RGBColor(0x00, 0xA8, 0x96)       # footer bar
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT  = RGBColor(0x1A, 0x1A, 0x1A)
SUBTITLE_GRAY = RGBColor(0x88, 0x88, 0x88)

# ── Dimensions ───────────────────────────────────────────────────────────────
SLIDE_W = Inches(10)
SLIDE_H = Inches(5.625)
MARGIN_LEFT = Inches(0.5)
CONTENT_W   = Inches(9.0)
HEADER_H    = Inches(0.8)
CONTENT_TOP = Inches(1.1)
FOOTER_H    = Inches(0.2)


def set_slide_bg(slide, color=BG_WHITE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header_bar(slide, title_text):
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, HEADER_H)
    header.fill.solid()
    header.fill.fore_color.rgb = HEADER_BG
    header.line.fill.background()

    # Small accent triangle/bar
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, HEADER_H, Inches(0.08), Inches(0.05))
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT
    accent.line.fill.background()

    txBox = slide.shapes.add_textbox(MARGIN_LEFT, Inches(0.18), CONTENT_W, Inches(0.55))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title_text
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = "Segoe UI"


def add_footer(slide):
    footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.425), SLIDE_W, FOOTER_H)
    footer.fill.solid()
    footer.fill.fore_color.rgb = FOOTER_BG
    footer.line.fill.background()


def add_content_box(slide, left, top, width, height):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    return tf


def add_bullet_para(tf, text, font_size=Pt(14), bold=False, color=DARK_TEXT,
                    space_before=Pt(4), space_after=Pt(4), bullet_char="\u2022"):
    if len(tf.paragraphs) == 1 and tf.paragraphs[0].text == "":
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    p.space_before = space_before
    p.space_after = space_after

    if bullet_char:
        run_b = p.add_run()
        run_b.text = f"{bullet_char}  "
        run_b.font.size = font_size
        run_b.font.bold = False
        run_b.font.color.rgb = ACCENT
        run_b.font.name = "Segoe UI"

    run = p.add_run()
    run.text = text
    run.font.size = font_size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Segoe UI"
    return p


def add_text_para(tf, text, font_size=Pt(14), bold=False, color=DARK_TEXT,
                  alignment=PP_ALIGN.LEFT, space_before=Pt(2), space_after=Pt(2)):
    if len(tf.paragraphs) == 1 and tf.paragraphs[0].text == "":
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.alignment = alignment
    p.space_before = space_before
    p.space_after = space_after

    run = p.add_run()
    run.text = text
    run.font.size = font_size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Segoe UI"
    return p


def add_label_value(tf, label, value, font_size=Pt(14), space_before=Pt(6), space_after=Pt(2)):
    if len(tf.paragraphs) == 1 and tf.paragraphs[0].text == "":
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    p.space_before = space_before
    p.space_after = space_after

    run_l = p.add_run()
    run_l.text = f"{label}: "
    run_l.font.size = font_size
    run_l.font.bold = True
    run_l.font.color.rgb = DARK_TEXT
    run_l.font.name = "Segoe UI"

    run_v = p.add_run()
    run_v.text = value
    run_v.font.size = font_size
    run_v.font.bold = False
    run_v.font.color.rgb = DARK_TEXT
    run_v.font.name = "Segoe UI"
    return p


# ══════════════════════════════════════════════════════════════════════════════
# BUILD PRESENTATION
# ══════════════════════════════════════════════════════════════════════════════

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
blank_layout = prs.slide_layouts[6]  # Blank layout

# ─── SLIDE 1 — TITLE ────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, TITLE_BG)

tf = add_content_box(slide, MARGIN_LEFT, Inches(1.6), CONTENT_W, Inches(1.2))
add_text_para(tf, "Material Management System", Pt(36), bold=True, color=WHITE,
              alignment=PP_ALIGN.CENTER, space_after=Pt(0))

tf2 = add_content_box(slide, MARGIN_LEFT, Inches(2.7), CONTENT_W, Inches(0.8))
add_text_para(tf2, "A Comprehensive Solution for Inventory & Resource Management",
              Pt(18), color=ACCENT, alignment=PP_ALIGN.CENTER)

tf3 = add_content_box(slide, MARGIN_LEFT, Inches(3.5), CONTENT_W, Inches(0.5))
add_text_para(tf3, "Karnataka Science College, Dharwad  |  Manvit G M  |  U02AB23S0350",
              Pt(11), color=RGBColor(0xCC, 0xEE, 0xE8), alignment=PP_ALIGN.CENTER)

footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(4.9), SLIDE_W, Inches(0.42))
footer.fill.solid()
footer.fill.fore_color.rgb = FOOTER_BG
footer.line.fill.background()

tf_f = add_content_box(slide, MARGIN_LEFT, Inches(4.95), CONTENT_W, Inches(0.3))
add_text_para(tf_f, "Material Management System", Pt(10), color=WHITE, alignment=PP_ALIGN.LEFT)

# ─── SLIDE 2 — INTRODUCTION ─────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Introduction")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, CONTENT_W, Inches(4.0))
add_label_value(tf, "Project Title", "Material Management System (MMS)", Pt(16), Pt(8), Pt(4))
add_label_value(tf, "Developed by", "Manvit G M", Pt(16), Pt(14), Pt(4))
add_label_value(tf, "UUCMS Registration Number", "U02AB23S0350", Pt(16), Pt(14), Pt(4))
add_label_value(tf, "Department", "Computer Science", Pt(16), Pt(14), Pt(4))
add_label_value(tf, "College", "Karnataka Science College, Dharwad", Pt(16), Pt(14), Pt(4))
add_label_value(tf, "Guide", "Faculty of Computer Science Department", Pt(16), Pt(14), Pt(4))

# ─── SLIDE 3 — ABSTRACT ─────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Abstract")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, CONTENT_W, Inches(4.0))
add_text_para(tf, "Overview:", Pt(16), bold=True, color=DARK_TEXT, space_after=Pt(8))
add_text_para(tf, (
    "Material management is a strategic function that integrates the planning, sourcing, "
    "and flow of goods to optimize production efficiency. By balancing procurement, inventory "
    "control, and warehousing, it ensures that the right materials are available at the right "
    "time while minimizing carrying costs."
), Pt(13), color=DARK_TEXT, space_after=Pt(10))
add_text_para(tf, (
    "Modern practices leverage ERP systems and Just-in-Time (JIT) strategies to reduce waste "
    "and improve cash flow. Ultimately, effective material management serves as a critical "
    "driver for profitability, as materials typically represent the largest portion of "
    "manufacturing expenses. It transforms the supply chain into a competitive advantage "
    "through data-driven logistics."
), Pt(13), color=DARK_TEXT, space_after=Pt(10))

# ─── SLIDE 4 — TABLE OF CONTENTS ────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Table of Contents")
add_footer(slide)

toc_items = [
    "Introduction & Project Overview",
    "Abstract & Problem Definition",
    "Existing System Analysis",
    "Proposed Solution & Architecture",
    "System Requirements & Technology Stack",
    "Data Flow Diagram (DFD)",
    "Entity-Relationship Diagram (ER)",
    "Implementation Plan & Methodology",
    "Testing & Quality Assurance",
    "Advantages & Key Benefits",
    "Conclusion",
    "Future Enhancements & References",
]

col_w = Inches(4.2)
tf_l = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, col_w, Inches(3.8))
for idx, item in enumerate(toc_items[:6]):
    add_bullet_para(tf_l, item, Pt(14), color=DARK_TEXT,
                    space_before=Pt(8), space_after=Pt(8), bullet_char=f"{idx+1}.")

tf_r = add_content_box(slide, Inches(5.2), CONTENT_TOP, col_w, Inches(3.8))
for idx, item in enumerate(toc_items[6:]):
    add_bullet_para(tf_r, item, Pt(14), color=DARK_TEXT,
                    space_before=Pt(8), space_after=Pt(8), bullet_char=f"{idx+7}.")

# ─── SLIDE 5 — OBJECTIVES ───────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Objectives & Goals")
add_footer(slide)

objectives = [
    "Maintain accurate digital records of materials and inventory",
    "Track stock levels in real-time to prevent shortages or overstocking",
    "Efficiently manage supplier details and purchase orders",
    "Provide secure and centralized storage of material data",
    "Generate comprehensive reports for stock, purchase, and usage analysis",
    "Reduce manual errors and improve operational efficiency",
    "Improve material traceability across departments",
    "Support faster decision-making with real-time analytics and reports",
    "Implement role-based access control for data security",
]

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, CONTENT_W, Inches(4.0))
for obj in objectives:
    add_bullet_para(tf, obj, Pt(13), color=DARK_TEXT, space_before=Pt(5), space_after=Pt(5))

# ─── SLIDE 6 — PROBLEM DEFINITION ───────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Problem Definition")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, CONTENT_W, Inches(4.0))
add_text_para(tf, "Current Challenges with Manual Systems:", Pt(15), bold=True,
              color=DARK_TEXT, space_after=Pt(10))

problems = [
    "Difficult to track materials accurately across departments",
    "Frequent errors in stock calculations and updates",
    "Delays in identifying material shortages leading to production halts",
    "Complex supplier and purchase order management",
    "Risk of data loss or damage in paper-based systems",
    "Lack of proper reporting and real-time analytics",
    "Lack of transparency and accountability in operations",
    "Time-consuming processes causing delays and customer dissatisfaction",
    "High risk of data duplication and unauthorised access",
]

for prob in problems:
    add_bullet_para(tf, prob, Pt(12), color=DARK_TEXT, space_before=Pt(4), space_after=Pt(4))

# ─── SLIDE 7 — EXISTING SYSTEM ──────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Existing System")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, CONTENT_W, Inches(4.0))
add_text_para(tf, "Current System Characteristics:", Pt(15), bold=True,
              color=DARK_TEXT, space_after=Pt(10))

existing = [
    "Paper-based manual record keeping for all material transactions",
    "Materials documented in physical registers and files",
    "Time-consuming manual searches for material information",
    "High frequency of errors in stock calculations and billing",
    "Risk of record loss, damage, or misuse",
    "No automated tracking of inventory usage and movements",
    "No analytical or reporting capabilities for informed decision-making",
    "Lack of transparency and accountability in material flow",
]

for item in existing:
    add_bullet_para(tf, item, Pt(13), color=DARK_TEXT, space_before=Pt(5), space_after=Pt(5))

# ─── SLIDE 8 — PROPOSED SYSTEM ──────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Proposed System")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, CONTENT_W, Inches(4.0))
add_text_para(tf, "Solution Overview:", Pt(15), bold=True, color=DARK_TEXT, space_after=Pt(6))
add_text_para(tf, (
    "A modern, computerized Material Management System providing end-to-end "
    "digital management of materials with the following features:"
), Pt(12), color=DARK_TEXT, space_after=Pt(8))

proposed = [
    "Digital records of all materials and inventory with real-time updates",
    "Quick and efficient access to inventory details and history",
    "Automated purchase and supplier management workflows",
    "Real-time report generation and analytics dashboard",
    "Improved accuracy and enhanced data security",
    "User-friendly interface with role-based access control",
    "Fully digitalised workflow accessible 24/7 from any internet-enabled device",
    "Complete audit trail and accountability for all transactions",
    "Automated notifications and reminders to reduce manual follow-up",
]

for item in proposed:
    add_bullet_para(tf, item, Pt(12), color=DARK_TEXT, space_before=Pt(3), space_after=Pt(3))

# ─── SLIDE 9 — FEASIBILITY STUDY ────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Feasibility Study")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, CONTENT_W, Inches(4.0))

add_text_para(tf, "Technical Feasibility:", Pt(14), bold=True, color=HEADER_BG, space_after=Pt(4))
add_text_para(tf, (
    "Django and MySQL provide a robust, scalable backend. Barcode integration uses "
    "open-source Python libraries. The system can be hosted on-premises or on cloud."
), Pt(12), color=DARK_TEXT, space_after=Pt(10))

add_text_para(tf, "Economic Feasibility:", Pt(14), bold=True, color=HEADER_BG, space_after=Pt(4))
add_text_para(tf, (
    "Reduces material losses, eliminates emergency procurement premiums, and cuts "
    "administrative labour, delivering a clear ROI within the first year."
), Pt(12), color=DARK_TEXT, space_after=Pt(10))

add_text_para(tf, "Operational Feasibility:", Pt(14), bold=True, color=HEADER_BG, space_after=Pt(4))
add_text_para(tf, (
    "Store managers and department heads with basic computer skills can use the system "
    "after a short training session. The workflow mirrors existing physical processes."
), Pt(12), color=DARK_TEXT, space_after=Pt(10))

add_text_para(tf, "Legal Feasibility:", Pt(14), bold=True, color=HEADER_BG, space_after=Pt(4))
add_text_para(tf, (
    "The system complies with standard data protection and procurement regulations "
    "through secure user authentication, role-based access control, and audit trails."
), Pt(12), color=DARK_TEXT, space_after=Pt(6))

# ─── SLIDE 10 — SYSTEM REQUIREMENTS ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "System Requirements")
add_footer(slide)

tf_hw = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, Inches(4.2), Inches(3.8))
add_text_para(tf_hw, "Hardware Requirements", Pt(15), bold=True, color=HEADER_BG, space_after=Pt(8))
for label, value in [
    ("Processor", "Intel Core i5 or above"),
    ("RAM", "4 GB minimum (8 GB recommended)"),
    ("Hard Disk", "500 GB HDD (1 TB SSD recommended)"),
    ("Network", "10 Mbps Broadband"),
    ("Display", "1366 x 768 (1920 x 1080 recommended)"),
]:
    add_label_value(tf_hw, label, value, Pt(12), Pt(6), Pt(2))

tf_sw = add_content_box(slide, Inches(5.2), CONTENT_TOP, Inches(4.3), Inches(3.8))
add_text_para(tf_sw, "Software Requirements", Pt(15), bold=True, color=HEADER_BG, space_after=Pt(8))
for label, value in [
    ("OS", "Windows 10 / 11"),
    ("Frontend", "HTML5, CSS3, JavaScript, Bootstrap 5"),
    ("Backend", "Python (Django Framework)"),
    ("Database", "MySQL / SQLite"),
    ("Web Server", "XAMPP / Django Dev Server"),
    ("IDE", "PyCharm / VS Code"),
    ("Version Control", "Git"),
]:
    add_label_value(tf_sw, label, value, Pt(12), Pt(6), Pt(2))

# ─── SLIDE 11 — DFD Level 0 ─────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Data Flow Diagram \u2014 Level 0 (Context Diagram)")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, Inches(4.0), Inches(3.8))
add_text_para(tf, "Context Diagram:", Pt(14), bold=True, color=HEADER_BG, space_after=Pt(6))
add_text_para(tf, (
    "The Level 0 DFD (Context Diagram) provides an abstraction view, "
    "showing the system as a single process with its relationship to "
    "external entities. It represents the entire system as a single bubble "
    "with input and output data indicated by incoming/outgoing arrows."
), Pt(11), color=DARK_TEXT, space_after=Pt(10))
add_text_para(tf, "Figure 3.3.1", Pt(10), bold=True, color=SUBTITLE_GRAY, space_before=Pt(6))

dfd0_path = r'c:\material\extracted_images\dfd_level0.jpeg'
if os.path.exists(dfd0_path):
    slide.shapes.add_picture(dfd0_path, Inches(5.0), Inches(1.1), Inches(4.5), Inches(3.6))

# ─── SLIDE 12 — DFD Level 1 ─────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Data Flow Diagram \u2014 Level 1")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, Inches(4.0), Inches(3.8))
add_text_para(tf, "Level 1 DFD:", Pt(14), bold=True, color=HEADER_BG, space_after=Pt(6))
add_text_para(tf, (
    "This level provides a more detailed view by breaking down the major "
    "processes into sub-processes. Each subprocess is depicted as a separate "
    "process with associated data flows and data stores."
), Pt(11), color=DARK_TEXT, space_after=Pt(8))

add_text_para(tf, "Key DFD Rules:", Pt(12), bold=True, color=HEADER_BG, space_after=Pt(4))
for rule in [
    "Each data store must have at least one data flow in and out",
    "Data stored must go through a process",
    "All processes must have unique names",
    "Inputs to a process should differ from outputs",
]:
    add_bullet_para(tf, rule, Pt(10), color=DARK_TEXT, space_before=Pt(2), space_after=Pt(2))

add_text_para(tf, "Figure 3.3.2", Pt(10), bold=True, color=SUBTITLE_GRAY, space_before=Pt(6))

dfd1_path = r'c:\material\extracted_images\dfd_level1.jpeg'
if os.path.exists(dfd1_path):
    slide.shapes.add_picture(dfd1_path, Inches(5.0), Inches(1.1), Inches(4.5), Inches(3.4))

# ─── SLIDE 13 — ER DIAGRAM ──────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Entity-Relationship Diagram (ER)")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, Inches(3.8), Inches(4.0))
add_text_para(tf, "ER Diagram Overview:", Pt(14), bold=True, color=HEADER_BG, space_after=Pt(6))
add_text_para(tf, (
    "An entity-relationship diagram is a data modeling technique that creates "
    "a graphical representation of entities and their relationships within "
    "the information system."
), Pt(11), color=DARK_TEXT, space_after=Pt(8))

add_text_para(tf, "Key Entities:", Pt(12), bold=True, color=HEADER_BG, space_after=Pt(4))
for ent in [
    "Users \u2014 User ID, Name, Email, Role, Dept",
    "Materials \u2014 Material ID, Name, Category, Price",
    "Suppliers \u2014 Supplier ID, Company, Contact",
    "Orders \u2014 Order ID, Date, Status, Amount",
    "Inventory \u2014 Stock ID, Qty, Location",
    "Transactions \u2014 Txn ID, Type, Date, Amount",
]:
    add_bullet_para(tf, ent, Pt(10), color=DARK_TEXT, space_before=Pt(2), space_after=Pt(2))

add_text_para(tf, "Figure 3.3.3", Pt(10), bold=True, color=SUBTITLE_GRAY, space_before=Pt(6))

er_path = r'c:\material\extracted_images\er_diagram.jpeg'
if os.path.exists(er_path):
    slide.shapes.add_picture(er_path, Inches(5.0), Inches(0.9), Inches(4.5), Inches(4.3))

# ─── SLIDE 14 — IMPLEMENTATION ──────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Implementation")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, Inches(4.2), Inches(4.0))
add_text_para(tf, "Development Phases:", Pt(15), bold=True, color=HEADER_BG, space_after=Pt(8))
for label, desc in [
    ("Phase 1", "Requirements Analysis & Planning"),
    ("Phase 2", "System Design & Architecture"),
    ("Phase 3", "Database Design & Development"),
    ("Phase 4", "Frontend Development & UI/UX"),
    ("Phase 5", "Backend Development & Integration"),
    ("Phase 6", "Testing & Quality Assurance"),
    ("Phase 7", "Deployment & User Training"),
]:
    add_label_value(tf, label, desc, Pt(12), Pt(6), Pt(3))

tf2 = add_content_box(slide, Inches(5.2), CONTENT_TOP, Inches(4.3), Inches(4.0))
add_text_para(tf2, "Languages & Tools Used:", Pt(15), bold=True, color=HEADER_BG, space_after=Pt(8))
for label, desc in [
    ("Python", "Primary backend language with Django framework"),
    ("HTML5", "Structure and content of web pages"),
    ("CSS3", "Styling and responsive design"),
    ("JavaScript", "Client-side interactivity and validation"),
    ("Bootstrap 5", "Responsive UI component framework"),
    ("MySQL", "Relational database management"),
]:
    add_label_value(tf2, label, desc, Pt(11), Pt(6), Pt(2))

# ─── SLIDE 15 — TESTING ─────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Testing & Quality Assurance")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, CONTENT_W, Inches(4.0))
add_text_para(tf, "Testing Methodology (V-Model):", Pt(14), bold=True, color=HEADER_BG, space_after=Pt(8))

for label, desc in [
    ("Unit Testing", "Individual functions tested in isolation using Django's TestCase class"),
    ("Integration Testing", "Modules tested in combination to verify data flow between components"),
    ("System Testing", "End-to-end testing of complete application workflows"),
    ("User Acceptance Testing", "Real users validate the system meets business requirements"),
]:
    add_label_value(tf, label, desc, Pt(12), Pt(8), Pt(4))

add_text_para(tf, "", Pt(6), space_before=Pt(6), space_after=Pt(2))
add_text_para(tf, "Sample Test Cases:", Pt(13), bold=True, color=HEADER_BG, space_after=Pt(4))

for tc in [
    "TC001 \u2014 Material Registration \u2192 Material added to inventory \u2713",
    "TC002 \u2014 Requisition Submission \u2192 Requisition routed to approver \u2713",
    "TC003 \u2014 Approval Workflow \u2192 Status changes to Approved \u2713",
    "TC004 \u2014 PO Generation \u2192 PO emailed to supplier \u2713",
    "TC005 \u2014 GRN Entry \u2192 Stock updated in real-time \u2713",
]:
    add_bullet_para(tf, tc, Pt(10), color=DARK_TEXT, space_before=Pt(2), space_after=Pt(2))

# ─── SLIDE 16 — ADVANTAGES ──────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Advantages & Key Benefits")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, CONTENT_W, Inches(4.0))

for adv in [
    "Simplifies material and inventory management with digital workflows",
    "Reduces manual errors and operational costs significantly",
    "Provides accurate, real-time stock information across departments",
    "Enables better decision-making through comprehensive analytics",
    "Enhances data security with role-based access control",
    "Integrates CAD designing for better material visualization",
    "Improves efficiency and reduces operational time",
    "Complete audit trail ensures accountability for all material movements",
    "Scalable architecture supports growing data volumes and user base",
    "Automated notifications reduce manual follow-up effort",
]:
    add_bullet_para(tf, adv, Pt(12), color=DARK_TEXT, space_before=Pt(4), space_after=Pt(4))

# ─── SLIDE 17 — CONCLUSION ──────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Conclusion")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, CONTENT_W, Inches(4.0))
add_text_para(tf, (
    "The Material Management System provides a comprehensive, automated solution for "
    "organisations to manage their materials efficiently."
), Pt(14), color=DARK_TEXT, space_after=Pt(12))

add_text_para(tf, "Key Outcomes:", Pt(14), bold=True, color=HEADER_BG, space_after=Pt(8))

for c in [
    "Real-time inventory visibility eliminates stockouts and overstocking",
    "Digital workflows speed up procurement and issuance processes",
    "Audit trails ensure full accountability for material movements",
    "Significant cost savings through reduced material wastage",
    "Better vendor negotiation from performance data analytics",
    "Lower administrative overhead with automated processes",
    "Future scope includes ERP integration, IoT sensors, RFID tracking, and mobile apps",
]:
    add_bullet_para(tf, c, Pt(12), color=DARK_TEXT, space_before=Pt(5), space_after=Pt(5))

# ─── SLIDE 18 — FUTURE ENHANCEMENTS ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "Future Enhancements")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, CONTENT_W, Inches(4.0))
add_text_para(tf, "Planned Enhancements:", Pt(15), bold=True, color=DARK_TEXT, space_after=Pt(8))

for label, desc in [
    ("Mobile Application", "Dedicated Android & iOS app for on-the-go access"),
    ("AI Integration", "Machine learning for predictive analytics & demand forecasting"),
    ("Cloud Deployment", "Migration to AWS/Azure for high availability & auto-scaling"),
    ("API Development", "RESTful API for third-party & government portal integration"),
    ("IoT Integration", "Automated inventory tracking with RFID & IoT sensors"),
    ("Advanced Reporting", "BI dashboards with interactive charts & drill-down capabilities"),
    ("Multi-language Support", "Regional language interfaces for wider accessibility"),
    ("ERP Integration", "Integration with Tally, SAP, and other ERP systems"),
]:
    add_label_value(tf, label, desc, Pt(12), Pt(6), Pt(3))

# ─── SLIDE 19 — REFERENCES ──────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, BG_WHITE)
add_header_bar(slide, "References")
add_footer(slide)

tf = add_content_box(slide, MARGIN_LEFT, CONTENT_TOP, CONTENT_W, Inches(4.0))
add_text_para(tf, "Technical References:", Pt(15), bold=True, color=DARK_TEXT, space_after=Pt(10))

for ref in [
    "Python Official Documentation \u2014 https://www.python.org/",
    "Django Documentation \u2014 https://docs.djangoproject.com/",
    "Bootstrap Documentation \u2014 https://getbootstrap.com/",
    "MySQL Documentation \u2014 https://dev.mysql.com/doc/",
    "SQLite Documentation \u2014 https://www.sqlite.org/",
    "HTML/CSS/JavaScript Standards \u2014 https://www.w3.org/",
    "XAMPP \u2014 https://www.apachefriends.org/",
    "Software Engineering Best Practices \u2014 IEEE Standards",
]:
    add_bullet_para(tf, ref, Pt(12), color=DARK_TEXT, space_before=Pt(5), space_after=Pt(5))

# ─── SLIDE 20 — THANK YOU ───────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, TITLE_BG)

tf = add_content_box(slide, MARGIN_LEFT, Inches(1.8), CONTENT_W, Inches(1.0))
add_text_para(tf, "Thank You!", Pt(40), bold=True, color=WHITE,
              alignment=PP_ALIGN.CENTER, space_after=Pt(0))

tf2 = add_content_box(slide, MARGIN_LEFT, Inches(2.8), CONTENT_W, Inches(0.6))
add_text_para(tf2, "Questions & Discussion", Pt(20), color=ACCENT,
              alignment=PP_ALIGN.CENTER)

tf3 = add_content_box(slide, MARGIN_LEFT, Inches(3.6), CONTENT_W, Inches(0.5))
add_text_para(tf3, "Material Management System  |  Manvit G M  |  Karnataka Science College, Dharwad",
              Pt(11), color=RGBColor(0xCC, 0xEE, 0xE8), alignment=PP_ALIGN.CENTER)

footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(4.9), SLIDE_W, Inches(0.42))
footer.fill.solid()
footer.fill.fore_color.rgb = FOOTER_BG
footer.line.fill.background()

tf_f = add_content_box(slide, MARGIN_LEFT, Inches(4.95), CONTENT_W, Inches(0.3))
add_text_para(tf_f, "Material Management System", Pt(10), color=WHITE, alignment=PP_ALIGN.LEFT)


# ── SAVE ─────────────────────────────────────────────────────────────────────
output_path = r'C:\Users\manvi\Downloads\Enhanced_MMS_Presentation_Final.pptx'
prs.save(output_path)
print(f"\u2705 Presentation saved to: {output_path}")
print(f"   Total slides: {len(prs.slides)}")
print(f"   File size: {os.path.getsize(output_path):,} bytes")
